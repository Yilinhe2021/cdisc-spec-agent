"""Write findings back into a copy of the spec workbook, next to the rows they concern."""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import SEVERITY_ORDER, QCResult, Spec

FILL = {"critical": "F8B4B4", "major": "FCD9A8", "minor": "FFF3B0", "info": "DCEBFA"}


def _note(fs) -> str:
    return "\n".join(f"[{f.severity.value.upper()}][{f.rule_id}] {f.message}" + (f"\n  -> {f.suggestion}" if f.suggestion else "")
                     for f in fs)


def annotate(spec_path: Path | None, spec: Spec, result: QCResult, out_path: Path) -> Path:
    findings = [f for f in result.findings if f.status != "dismissed"]
    if spec_path and spec_path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(spec_path)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Variables"
        ws.append(["Dataset", "Variable", "Label", "Type", "Length", "Origin", "Derivation", "Codelist", "Core"])
        for v in spec.variables:
            ws.append([v.dataset, v.variable, v.label, v.type, v.length, v.origin, v.derivation, v.codelist, v.core])
        spec.sheet_names["variables"] = "Variables"

    def add_cols(sheet_name: str, by_row: dict[int, list]):
        ws = wb[sheet_name]
        c0 = ws.max_column + 1
        for j, h in enumerate(("QC Severity", "QC Findings")):
            cell = ws.cell(row=1, column=c0 + j, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.column_dimensions[get_column_letter(c0)].width = 12
        ws.column_dimensions[get_column_letter(c0 + 1)].width = 80
        for row, fs in by_row.items():
            worst = min(fs, key=lambda f: SEVERITY_ORDER[f.severity]).severity.value
            ws.cell(row=row, column=c0, value=worst).fill = PatternFill("solid", fgColor=FILL[worst])
            c = ws.cell(row=row, column=c0 + 1, value=_note(fs))
            c.alignment = Alignment(wrap_text=True, vertical="top")

    var_rows: dict[int, list] = defaultdict(list)
    ds_rows: dict[int, list] = defaultdict(list)
    for f in findings:
        if f.variable and f.row:
            var_rows[f.row].append(f)
        elif f.row and "datasets" in spec.sheet_names and f.dataset != "(codelists)":
            ds_rows[f.row].append(f)
    add_cols(spec.sheet_names["variables"], var_rows)
    if ds_rows:
        add_cols(spec.sheet_names["datasets"], ds_rows)

    ws = wb.create_sheet("QC_Findings", 0)
    ws.append(["#", "Severity", "Category", "Dataset", "Variable", "Row", "Rule", "Agent", "Finding", "Suggestion",
               "Evidence", "Confidence", "Status"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for i, f in enumerate(findings, 1):
        ws.append([i, f.severity.value, f.category.value, f.dataset, f.variable, f.row, f.rule_id, f.agent,
                   f.message, f.suggestion, "; ".join(f"{e.source} > {e.section}" for e in f.evidence),
                   round(f.confidence, 2), f.status])
        ws.cell(row=i + 1, column=2).fill = PatternFill("solid", fgColor=FILL[f.severity.value])
    for col, w in zip("ABCDEFGHIJKLM", (5, 10, 13, 10, 11, 6, 14, 18, 70, 50, 35, 10, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
